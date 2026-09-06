#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from project_manifest import DOCUMENTS, load_manifest, manifest_path, save_manifest, outputs_for, source_fingerprint
from applicant_policy import load_company_config, subject_issues


MEDIA_SUFFIXES = {".mp4", ".mov", ".avi", ".mkv", ".m4v", ".webm", ".mp3", ".wav", ".m4a", ".aac", ".flac"}
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".heic", ".webp", ".tif", ".tiff"}
DOC_SUFFIXES = {".pdf", ".doc", ".docx", ".txt", ".md", ".csv", ".rtf", ".xls", ".xlsx"}
SKIP_DIRS = {".git", "__pycache__", ".venv", "venv", "node_modules"}


def load_manifest_optional(project_dir: Path) -> dict[str, Any] | None:
    path = manifest_path(project_dir)
    if not path.exists():
        return None
    return load_manifest(path)


def rel(path: Path, roots: list[Path]) -> str:
    resolved = path.resolve()
    for root in roots:
        try:
            return str(resolved.relative_to(root.resolve()))
        except ValueError:
            continue
    return str(resolved)


def human_size(size: int) -> str:
    units = ["B", "KB", "MB", "GB"]
    value = float(size)
    for unit in units:
        if value < 1024 or unit == units[-1]:
            if unit == "B":
                return f"{int(value)}{unit}"
            return f"{value:.1f}{unit}"
        value /= 1024
    return f"{size}B"


def should_skip(path: Path) -> bool:
    return any(part in SKIP_DIRS for part in path.parts)


def classify(path: Path) -> str:
    name = path.name.lower()
    suffix = path.suffix.lower()
    if suffix in {".mp4", ".mov", ".avi", ".mkv", ".m4v", ".webm"}:
        return "video"
    if suffix in {".mp3", ".wav", ".m4a", ".aac", ".flac"}:
        return "audio"
    if "歌词" in name or "lyric" in name:
        return "lyrics"
    if "节目" in name or "曲目" in name or "setlist" in name:
        return "program"
    if "身份证" in name or "证件" in name or "护照" in name or suffix in IMAGE_SUFFIXES:
        return "identity"
    if re.match(r"^(0[0-9]|10)[_、 -]", path.name):
        return "final_or_template"
    if suffix in DOC_SUFFIXES:
        return "documents"
    if suffix in MEDIA_SUFFIXES:
        return "media"
    return "other"


def iter_candidate_files(roots: list[Path]) -> list[Path]:
    files: list[Path] = []
    seen = set()
    for root in roots:
        if not root.exists():
            continue
        for path in sorted(root.rglob("*")):
            if path.is_file() and not should_skip(path) and path.resolve() not in seen:
                files.append(path)
                seen.add(path.resolve())
    return files


def file_record(path: Path, roots: list[Path]) -> dict[str, Any]:
    stat = path.stat()
    return {
        "path": rel(path, roots),
        "absolute_path": str(path.resolve()),
        "name": path.name,
        "bucket": classify(path),
        "suffix": path.suffix.lower(),
        "size": stat.st_size,
        "size_human": human_size(stat.st_size),
        "mtime": int(stat.st_mtime),
        "mtime_text": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
    }


def normalize_text(text: str, limit: int) -> str:
    text = re.sub(r"\s+", " ", text).strip()
    if len(text) <= limit:
        return text
    return text[:limit].rstrip() + "..."


def extract_text_preview(path: Path, limit: int) -> str | None:
    suffix = path.suffix.lower()
    try:
        if suffix in {".txt", ".md", ".csv"}:
            return normalize_text(path.read_text(encoding="utf-8", errors="ignore"), limit)
        if suffix == ".docx":
            from docx import Document  # type: ignore

            document = Document(str(path))
            parts: list[str] = []
            parts.extend(p.text for p in document.paragraphs if p.text.strip())
            for table in document.tables[:3]:
                for row in table.rows[:12]:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        parts.append(row_text)
            return normalize_text("\n".join(parts), limit)
        if suffix == ".pdf":
            from pypdf import PdfReader  # type: ignore

            reader = PdfReader(str(path))
            parts = []
            for page in reader.pages[:2]:
                parts.append(page.extract_text() or "")
            return normalize_text("\n".join(parts), limit)
    except Exception as exc:
        return f"[text preview failed: {exc.__class__.__name__}]"
    return None


def final_output_status(project_dir: Path, manifest: dict[str, Any] | None) -> list[dict[str, Any]]:
    final_dir = Path(manifest["project"]["final_dir"]) if manifest else project_dir / "03_最终提交材料"
    rows: list[dict[str, Any]] = []
    for seq, meta in DOCUMENTS.items():
        outputs = outputs_for(final_dir, seq)
        rows.append(
            {
                "seq": seq,
                "name": meta["name"],
                "status": "present_unverified" if outputs else ("missing" if (manifest or {}).get("documents", {}).get(seq, meta).get("required") else "not_required_or_conditional"),
                "outputs": [path.name for path in outputs[:12]],
                "output_count": len(outputs),
            }
        )
    return rows


def manifest_summary(manifest: dict[str, Any] | None) -> dict[str, Any]:
    if not manifest:
        return {"exists": False}
    people = manifest.get("people", {})
    program = manifest.get("program", {})
    workflow = manifest.get("workflow", {})
    event = manifest.get("event", {})
    subject = manifest.get("application_subject", {})
    return {
        "exists": True,
        "approval_type": workflow.get("approval_type", ""),
        "locked": workflow.get("locked", {}),
        "application_subject": subject.get("company", subject.get("short_name", "")),
        "event_name": event.get("name_with_translation") or event.get("name", ""),
        "event_time": event.get("time_for_forms") or event.get("time_raw", ""),
        "performer_count": len(people.get("performers", [])),
        "excluded_staff_count": len(people.get("staff_excluded", [])),
        "id_warning_count": len(people.get("id_warnings", [])),
        "program_count": len(program.get("items", [])),
        "translation_count": len(program.get("translations", {})),
        "sensitive_edit_count": len(program.get("sensitive_edits", [])),
    }


def build_pack(project_dir: Path, source_dirs: list[Path], args: argparse.Namespace) -> dict[str, Any]:
    roots = [project_dir, *source_dirs]
    manifest = load_manifest_optional(project_dir)
    policy_issues = []
    if manifest:
        try:
            config, _ = load_company_config(args.config)
            policy_issues = subject_issues(manifest, config)
        except (ValueError, OSError) as exc:
            policy_issues = [str(exc)]
    mrz_scan_note = ""
    if args.scan_mrz:
        mrz_scan_note = run_mrz_scan(project_dir, source_dirs or [project_dir / "01_原始资料"], args.mrz_all_images)
    cn_id_scan_note = ""
    if args.scan_cn_id:
        cn_id_scan_note = run_cn_id_scan(project_dir, source_dirs or [project_dir / "01_原始资料"], args.cn_id_all_images)
    scan_roots = source_dirs or [project_dir / "01_原始资料", project_dir / "03_最终提交材料", project_dir / "04_缺失资料与提醒"]
    records = [file_record(path, roots) for path in iter_candidate_files(scan_roots)]
    records.sort(key=lambda item: (item["bucket"], item["path"]))

    by_bucket: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        by_bucket[record["bucket"]].append(record)

    text_previews: list[dict[str, str]] = []
    if args.extract_text:
        text_candidates = [
            Path(record["absolute_path"])
            for record in records
            if record["suffix"] in {".txt", ".md", ".csv", ".docx", ".pdf"}
        ]
        for path in text_candidates[: args.max_text_files]:
            preview = extract_text_preview(path, args.text_chars)
            if preview:
                text_previews.append({"path": rel(path, roots), "preview": preview})

    compact_buckets = {
        bucket: {
            "count": len(items),
            "total_size": sum(item["size"] for item in items),
            "total_size_human": human_size(sum(item["size"] for item in items)),
            "sample": items[: args.max_files_per_bucket],
        }
        for bucket, items in sorted(by_bucket.items())
    }

    pack: dict[str, Any] = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "project_dir": str(project_dir.resolve()),
        "source_dirs": [str(path.resolve()) for path in source_dirs],
        "manifest": manifest_summary(manifest),
        "applicant_policy_issues": policy_issues,
        "final_outputs": final_output_status(project_dir, manifest),
        "passport_mrz": load_passport_mrz_results(project_dir),
        "passport_mrz_scan_note": mrz_scan_note,
        "cn_id_ocr": load_cn_id_ocr_results(project_dir),
        "cn_id_ocr_scan_note": cn_id_scan_note,
        "source_inventory": compact_buckets,
        "text_previews": text_previews,
        "token_saving_notes": [
            "先读本 context_pack，再决定是否打开原始文件。",
            "不要把完整歌词、PDF 文本或长文件清单直接贴进上下文；需要时用脚本抽取局部预览。",
            "如果只改某一份文件，用 generate_domestic_package.py --only <序号> 局部重跑。",
            "视频先看压缩缓存和压缩报告，不重复读取或压缩未变化的大视频。",
            "本地 MRZ 提供有来源的候选字段；缺签发日起等字段仍需看原图，不把校验位通过当作全部字段正确。",
            "国内身份证先跑 scan_cn_id_ocr.py，本地 OCR 成功时只让模型做少量视觉抽检。",
        ],
        "high_token_steps": [
            "把完整素材目录树、长文件名列表和全部路径直接发给模型。",
            "让模型逐个打开 DOCX/PDF/歌词全文，尤其是多曲歌词和外文翻译材料。",
            "让模型视觉检查大量身份证、护照、PDF 页面截图；应先生成 contact sheet 或抽检集。",
            "重复读取完整 项目主档.json、执行计划、QA 报告，而不是读 compact context pack。",
            "把视频压缩日志、ffmpeg 输出或全部文件匹配过程贴进上下文。",
            "在模型里反复做可确定的校验，如身份证校验、MRZ 校验、文件存在性检查和视频参数检查。",
        ],
    }
    if manifest is not None:
        manifest.setdefault("cache", {})["last_agent_context_pack"] = str(
            project_dir / "02_生成中间文件" / "agent_context" / "context_pack.md"
        )
        save_manifest(manifest, manifest_path(project_dir))
    return pack


def run_mrz_scan(project_dir: Path, source_dirs: list[Path], all_images: bool) -> str:
    command = [
        sys.executable,
        str(Path(__file__).resolve().parent / "scan_passport_mrz.py"),
        "--project-dir",
        str(project_dir),
    ]
    for source_dir in source_dirs:
        command.extend(["--source-dir", str(source_dir)])
    if all_images:
        command.append("--all-images")
    result = subprocess.run(command, text=True, capture_output=True)
    if result.returncode == 0:
        return "scan_passport_mrz.py completed"
    return f"scan_passport_mrz.py failed: {result.stderr.strip() or result.stdout.strip()}"


def run_cn_id_scan(project_dir: Path, source_dirs: list[Path], all_images: bool) -> str:
    command = [
        sys.executable,
        str(Path(__file__).resolve().parent / "scan_cn_id_ocr.py"),
        "--project-dir",
        str(project_dir),
    ]
    for source_dir in source_dirs:
        command.extend(["--source-dir", str(source_dir)])
    if all_images:
        command.append("--all-images")
    result = subprocess.run(command, text=True, capture_output=True)
    if result.returncode == 0:
        return "scan_cn_id_ocr.py completed"
    return f"scan_cn_id_ocr.py failed: {result.stderr.strip() or result.stdout.strip()}"


def verify_extraction_sources(results: list[dict[str, Any]]) -> list[dict[str, Any]]:
    for item in results:
        source = Path(item["source_file"]) if item.get("source_file") else None
        saved = item.get("source_fingerprint")
        current = source_fingerprint(source) if source and source.is_file() else None
        if not saved or current != saved:
            item["status"] = "stale_or_unverified"
            item["needs_manual_review"] = True
            item["source_warning"] = "源文件缺失、变化或旧识别结果无指纹；先复核或重扫。"
    return results


def load_passport_mrz_results(project_dir: Path) -> dict[str, Any]:
    path = project_dir / "02_生成中间文件" / "passport_mrz" / "passport_mrz_results.json"
    if not path.exists():
        return {"exists": False, "results": []}
    try:
        results = json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        return {"exists": True, "error": f"{exc.__class__.__name__}: {exc}", "results": []}
    results = verify_extraction_sources(results)
    return {
        "exists": True,
        "path": str(path),
        "total": len(results),
        "ok": sum(1 for item in results if item.get("status") == "ok"),
        "needs_manual_review": sum(1 for item in results if item.get("needs_manual_review")),
        "results": results,
    }


def load_cn_id_ocr_results(project_dir: Path) -> dict[str, Any]:
    path = project_dir / "02_生成中间文件" / "cn_id_ocr" / "cn_id_ocr_results.json"
    if not path.exists():
        return {"exists": False, "results": []}
    try:
        results = json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        return {"exists": True, "error": f"{exc.__class__.__name__}: {exc}", "results": []}
    results = verify_extraction_sources(results)
    return {
        "exists": True,
        "path": str(path),
        "total": len(results),
        "ok": sum(1 for item in results if item.get("status") == "ok"),
        "needs_manual_review": sum(1 for item in results if item.get("needs_manual_review")),
        "renewal_warning": sum(1 for item in results if item.get("parsed", {}).get("renewal_warning")),
        "results": results,
    }


def render_markdown(pack: dict[str, Any]) -> str:
    manifest = pack["manifest"]
    lines = [
        "# Agent Context Pack",
        "",
        f"- generated_at: `{pack['generated_at']}`",
        f"- project_dir: `{pack['project_dir']}`",
        f"- manifest_exists: `{manifest.get('exists')}`",
    ]
    if pack.get("applicant_policy_issues"):
        lines.extend("- 主体政策阻断：" + issue for issue in pack["applicant_policy_issues"])
    if manifest.get("exists"):
        lines.extend(
            [
                f"- approval_type: {manifest.get('approval_type')}",
                f"- event: {manifest.get('event_name')}",
                f"- time: {manifest.get('event_time')}",
                f"- subject: {manifest.get('application_subject')}",
                f"- locked: `{json.dumps(manifest.get('locked', {}), ensure_ascii=False)}`",
                f"- performers: {manifest.get('performer_count')} / program_items: {manifest.get('program_count')}",
                f"- warnings: id={manifest.get('id_warning_count')} sensitive={manifest.get('sensitive_edit_count')}",
            ]
        )
    lines.extend(["", "## Final Outputs", ""])
    for row in pack["final_outputs"]:
        outputs = ", ".join(row["outputs"]) if row["outputs"] else "-"
        lines.append(f"- `{row['seq']}` {row['name']}: {row['status']} ({row['output_count']}) {outputs}")

    mrz = pack.get("passport_mrz", {})
    lines.extend(["", "## Passport / Permit MRZ", ""])
    if pack.get("passport_mrz_scan_note"):
        lines.append(f"- scan_note: {pack['passport_mrz_scan_note']}")
    if mrz.get("exists"):
        lines.append(f"- total: {mrz.get('total')} / ok: {mrz.get('ok')} / needs_manual_review: {mrz.get('needs_manual_review')}")
        for item in mrz.get("results", [])[:12]:
            status = item.get("status")
            name = Path(item.get("source_file") or "MRZ lines").name
            number = item.get("number") or ""
            expiry = item.get("expiration_date") or ""
            country = item.get("country_zh") or item.get("nationality") or item.get("country") or "待确认"
            approval_region = item.get("region_zh_for_approval") or country
            approval_doc = item.get("document_type_zh_for_approval") or "证件类型待确认"
            lines.append(f"- `{status}` {name}: {number} / {expiry} / {approval_region} / {approval_doc}")
            if item.get("error"):
                lines.append(f"  error: {item['error']}")
    else:
        lines.append("- 未找到本地 MRZ 结果。涉外项目先运行 `scan_passport_mrz.py` 或本脚本 `--scan-mrz`。")

    cn_id = pack.get("cn_id_ocr", {})
    lines.extend(["", "## Chinese ID OCR", ""])
    if pack.get("cn_id_ocr_scan_note"):
        lines.append(f"- scan_note: {pack['cn_id_ocr_scan_note']}")
    if cn_id.get("exists"):
        lines.append(
            f"- total: {cn_id.get('total')} / ok: {cn_id.get('ok')} / "
            f"needs_manual_review: {cn_id.get('needs_manual_review')} / renewal_warning: {cn_id.get('renewal_warning')}"
        )
        for item in cn_id.get("results", [])[:12]:
            parsed = item.get("parsed", {})
            status = item.get("status")
            name = Path(item.get("source_file") or "ID image").name
            person = parsed.get("name") or "待复核"
            expiry = parsed.get("valid_to") or "未识别"
            if parsed.get("side") == "back":
                identity_note = f"authority={parsed.get('issuing_authority') or '待复核'}"
            else:
                identity_note = f"id_valid={parsed.get('id_validation', {}).get('valid')}"
            lines.append(f"- `{status}` {name}: {person} / {identity_note} / expiry={expiry}")
            if parsed.get("renewal_warning"):
                lines.append(f"  warning: {person}的身份证需要换新。")
            if item.get("error"):
                lines.append(f"  error: {item['error']}")
    else:
        lines.append("- 未找到本地身份证 OCR 结果。国内项目可先运行 `scan_cn_id_ocr.py` 或本脚本 `--scan-cn-id`。")

    lines.extend(["", "## Source Inventory", ""])
    for bucket, info in pack["source_inventory"].items():
        lines.append(f"### {bucket} ({info['count']}, {info['total_size_human']})")
        for item in info["sample"]:
            lines.append(f"- `{item['path']}` {item['size_human']} {item['mtime_text']}")
        lines.append("")

    if pack["text_previews"]:
        lines.extend(["## Text Previews", ""])
        for preview in pack["text_previews"]:
            lines.append(f"### {preview['path']}")
            lines.append(preview["preview"])
            lines.append("")

    lines.extend(["## Token Saving Notes", ""])
    lines.extend(f"- {item}" for item in pack["token_saving_notes"])
    lines.extend(["", "## High Token Steps", ""])
    lines.extend(f"- {item}" for item in pack["high_token_steps"])
    lines.append("")
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(description="生成给 agent 使用的轻量上下文包，减少反复读大目录和长文档。")
    parser.add_argument("--project-dir", required=True, type=Path)
    parser.add_argument("--config", type=Path)
    parser.add_argument("--source-dir", action="append", default=[], type=Path)
    parser.add_argument("--extract-text", action="store_true", help="为 docx/pdf/txt/csv 抽取短文本预览。")
    parser.add_argument("--scan-mrz", action="store_true", help="先本地扫描护照/港澳台证件 MRZ，并写入 context pack。")
    parser.add_argument("--mrz-all-images", action="store_true", help="MRZ 扫描时处理全部图片/PDF，而不只处理护照/证件关键词文件。")
    parser.add_argument("--scan-cn-id", action="store_true", help="先本地扫描中国居民身份证 OCR，并写入 context pack。")
    parser.add_argument("--cn-id-all-images", action="store_true", help="身份证 OCR 时处理全部图片/PDF，而不只处理身份证关键词文件。")
    parser.add_argument("--max-text-files", type=int, default=25)
    parser.add_argument("--text-chars", type=int, default=600)
    parser.add_argument("--max-files-per-bucket", type=int, default=24)
    parser.add_argument("--output-dir", type=Path)
    args = parser.parse_args()

    project_dir = args.project_dir.resolve()
    output_dir = args.output_dir or project_dir / "02_生成中间文件" / "agent_context"
    output_dir.mkdir(parents=True, exist_ok=True)

    pack = build_pack(project_dir, [path.resolve() for path in args.source_dir], args)
    json_path = output_dir / "context_pack.json"
    md_path = output_dir / "context_pack.md"
    json_path.write_text(json.dumps(pack, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    md_path.write_text(render_markdown(pack), encoding="utf-8")
    print(md_path)
    print(json_path)


if __name__ == "__main__":
    main()
