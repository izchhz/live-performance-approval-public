#!/usr/bin/env python3
from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt


def add_run(paragraph, text: str, *, size: float = 14, bold: bool = False, underline: bool = False):
    run = paragraph.add_run(text)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold
    run.underline = underline
    return run


def add_paragraph(doc, *, before: float = 0, after: float = 0, line: float = 1.3, indent: float = 0):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    if indent:
        paragraph.paragraph_format.first_line_indent = Pt(indent)
    return paragraph


def main() -> None:
    parser = argparse.ArgumentParser(description="生成不包含个人隐私的授权委托书空白 Word 模板。")
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(27)
    section.bottom_margin = Mm(22)
    section.left_margin = Mm(30)
    section.right_margin = Mm(30)

    paragraph = add_paragraph(doc, after=20)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(paragraph, "授权委托书", size=24, bold=True)

    fields = [
        ("委托人名称（姓名）：", "{{委托人名称}}"),
        ("证件号码：", "{{委托人证件号码}}"),
        ("住所（住址）：", "{{委托人地址}}"),
        ("联系电话：", "{{委托人联系电话}}"),
        ("受委托人姓名：", "{{受委托人姓名}}"),
        ("身份证件号码：", "{{受委托人身份证件号码}}"),
        ("工作单位和职务：", "{{受委托人工作单位和职务}}"),
        ("住       址：", "{{受委托人地址}}"),
        ("联系电话：", "{{受委托人联系电话}}"),
    ]
    for label, value in fields:
        paragraph = add_paragraph(doc, after=4, line=1.15)
        add_run(paragraph, label)
        add_run(paragraph, "  ")
        add_run(paragraph, value, underline=True)

    paragraph = add_paragraph(doc, before=6, after=2, line=1.65, indent=28)
    add_run(paragraph, "现委托 ")
    add_run(paragraph, "{{受委托人姓名}}", underline=True)
    add_run(paragraph, " 同志，作为我（单位）在办理 “")
    add_run(paragraph, "{{演出名称}}", underline=True)
    add_run(paragraph, "” 事项的委托代理人。委托权限为特别授权，包括代为送交申请材料、代为接受调查询问、代为申请回避、代为陈述申辩、代为要求听证、代签各种法律文书等委托人所具有的所有权限。")

    paragraph = add_paragraph(doc, before=3, after=14, line=1.4, indent=28)
    add_run(paragraph, "委托期限自 ")
    add_run(paragraph, "{{生成日期}}", underline=True)
    add_run(paragraph, " 至 ")
    add_run(paragraph, "{{当年12月31日}}", underline=True)
    add_run(paragraph, "。")

    paragraph = add_paragraph(doc, after=6)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(paragraph, "委托人签名（盖章）")

    paragraph = add_paragraph(doc)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(paragraph, "{{生成日期}}")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)


if __name__ == "__main__":
    main()
